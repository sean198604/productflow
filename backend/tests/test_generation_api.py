from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from PIL import Image

from app.services.renderers import PptxRenderer, RenderRequest
from app.services.renderers.base import placeholder_png
from tests.conftest import IdentityFixture


def _login(client: TestClient, identifier: str, password: str) -> dict[str, str]:
    response = client.post(
        "/api/v1/auth/login",
        json={"identifier": identifier, "password": password},
    )
    assert response.status_code == 200, response.text
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def _xlsx_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Quotation"
    sheet["A1"] = "Product"
    sheet["A2"] = "Price"
    sheet["A3"] = "Customer"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _pptx_template() -> bytes:
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr(
            "[Content_Types].xml",
            """<?xml version="1.0" encoding="UTF-8"?>
            <Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
              <Default Extension="xml" ContentType="application/xml"/>
              <Default Extension="png" ContentType="image/png"/>
            </Types>""",
        )
        archive.writestr(
            "ppt/presentation.xml",
            """<?xml version="1.0" encoding="UTF-8"?>
            <p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
              <p:sldSz cx="12192000" cy="6858000"/>
            </p:presentation>""",
        )
        archive.writestr(
            "ppt/slides/slide1.xml",
            """<?xml version="1.0" encoding="UTF-8"?>
            <p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
              xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
              xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
              <p:cSld><p:spTree><p:nvGrpSpPr/><p:grpSpPr/>
                <p:sp><p:nvSpPr><p:cNvPr id="2" name="Product title"/></p:nvSpPr>
                  <p:txBody><a:p><a:r><a:t>Old title</a:t></a:r>
                    <a:r><a:t>Old suffix</a:t></a:r></a:p></p:txBody>
                </p:sp>
                <p:pic><p:nvPicPr><p:cNvPr id="3" name="Product image"/></p:nvPicPr>
                  <p:blipFill><a:blip r:embed="rId1"/>
                    <a:stretch><a:fillRect/></a:stretch></p:blipFill>
                  <p:spPr><a:xfrm><a:off x="100" y="100"/>
                    <a:ext cx="4000000" cy="3000000"/></a:xfrm></p:spPr>
                </p:pic>
              </p:spTree></p:cSld>
            </p:sld>""",
        )
        archive.writestr(
            "ppt/slides/_rels/slide1.xml.rels",
            """<?xml version="1.0" encoding="UTF-8"?>
            <Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
              <Relationship Id="rId1"
                Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
                Target="../media/image1.png"/>
            </Relationships>""",
        )
        archive.writestr("ppt/media/image1.png", placeholder_png(20, 20))
    return output.getvalue()


def test_customer_product_set_and_xlsx_generation_snapshot(
    client: TestClient,
    identity_fixture: IdentityFixture,
) -> None:
    headers = _login(
        client, identity_fixture.owner_username, identity_fixture.password
    )
    product_response = client.post(
        "/api/v1/products",
        headers=headers,
        json={
            "sku": "GEN-001",
            "product_name": "=FORMULA-MUST-STAY-TEXT",
            "status": "active",
            "custom_fields": {
                "selling_price": "12.5",
                "supplier_cost": "8.25",
            },
        },
    )
    assert product_response.status_code == 201, product_response.text
    product_id = product_response.json()["id"]

    template_response = client.post(
        "/api/v1/output-templates",
        headers=headers,
        data={"name": "Generation XLSX"},
        files={"file": ("generation.xlsx", _xlsx_template())},
    )
    assert template_response.status_code == 201, template_response.text
    template = template_response.json()
    version = template["current_version"]
    mapping_response = client.put(
        f"/api/v1/output-templates/versions/{version['id']}/mapping",
        headers=headers,
        json={
            "version": "1.0",
            "template_sha256": version["template_sha256"],
            "bindings": [
                {
                    "object_key": "xlsx:s1:cell:A1",
                    "source": "product_name",
                    "product_slot": 1,
                },
                {
                    "object_key": "xlsx:s1:cell:A2",
                    "source": "selling_price",
                    "formatter": "currency",
                    "product_slot": 1,
                },
                {
                    "object_key": "xlsx:s1:cell:A3",
                    "source": "customer.name",
                },
            ],
        },
    )
    assert mapping_response.status_code == 200, mapping_response.text

    customer_response = client.post(
        "/api/v1/customers",
        headers=headers,
        json={
            "name": "Generation Customer",
            "code": "GEN-CUSTOMER",
            "default_xlsx_template_id": template["id"],
            "settings": {
                "locale": "en-US",
                "currency": "USD",
                "timezone": "Asia/Shanghai",
                "settings": {},
            },
        },
    )
    assert customer_response.status_code == 201, customer_response.text
    customer = customer_response.json()
    assert customer["default_xlsx_template_id"] == template["id"]
    assert customer["template_bindings"][0]["is_default"] is True

    set_response = client.post(
        "/api/v1/product-sets",
        headers=headers,
        json={
            "name": "Generation Set",
            "customer_id": customer["id"],
            "product_ids": [product_id],
        },
    )
    assert set_response.status_code == 201, set_response.text
    product_set = set_response.json()
    assert product_set["items"][0]["sort_order"] == 0

    generation_response = client.post(
        "/api/v1/generation-tasks",
        headers=headers,
        json={
            "name": "Safe generated quote",
            "customer_id": customer["id"],
            "product_set_id": product_set["id"],
            "output_template_version_id": version["id"],
            "output_parameters": {"currency_symbol": "$"},
        },
    )
    assert generation_response.status_code == 201, generation_response.text
    task = generation_response.json()
    assert task["status"] == "completed", task
    assert task["download_url"]
    snapshot_fields = task["product_snapshot"][0]["fields"]
    assert snapshot_fields["product_name"] == "=FORMULA-MUST-STAY-TEXT"
    assert snapshot_fields["selling_price"] == "12.5"
    assert "supplier_cost" not in snapshot_fields
    assert task["template_snapshot"]["template_sha256"] == version["template_sha256"]

    download = client.get(task["download_url"], headers=headers)
    assert download.status_code == 200, download.text
    generated = load_workbook(BytesIO(download.content), data_only=False)
    sheet = generated["Quotation"]
    assert sheet["A1"].value == "=FORMULA-MUST-STAY-TEXT"
    assert sheet["A1"].data_type == "s"
    assert sheet["A2"].value == 12.5
    assert sheet["A2"].number_format == '"$"#,##0.00'
    assert sheet["A3"].value == "Generation Customer"
    generated.close()

    protected_product = client.delete(
        f"/api/v1/products/{product_id}", headers=headers
    )
    assert protected_product.status_code == 409
    assert "历史生成任务" in protected_product.json()["message"]
    protected_template = client.delete(
        f"/api/v1/output-templates/{template['id']}", headers=headers
    )
    assert protected_template.status_code == 409
    assert "历史生成任务" in protected_template.json()["message"]


def test_generation_rejects_more_products_than_mapped_slots(
    client: TestClient,
    identity_fixture: IdentityFixture,
) -> None:
    headers = _login(
        client, identity_fixture.owner_username, identity_fixture.password
    )
    product_ids: list[str] = []
    for suffix in ("A", "B"):
        response = client.post(
            "/api/v1/products",
            headers=headers,
            json={
                "sku": f"SLOT-{suffix}",
                "product_name": f"Slot Product {suffix}",
                "status": "active",
            },
        )
        assert response.status_code == 201, response.text
        product_ids.append(response.json()["id"])
    template_response = client.post(
        "/api/v1/output-templates",
        headers=headers,
        data={"name": "One Slot XLSX"},
        files={"file": ("one-slot.xlsx", _xlsx_template())},
    )
    version = template_response.json()["current_version"]
    saved = client.put(
        f"/api/v1/output-templates/versions/{version['id']}/mapping",
        headers=headers,
        json={
            "version": "1.0",
            "template_sha256": version["template_sha256"],
            "bindings": [
                {
                    "object_key": "xlsx:s1:cell:A1",
                    "source": "product_name",
                    "product_slot": 1,
                }
            ],
        },
    )
    assert saved.status_code == 200, saved.text
    customer = client.post(
        "/api/v1/customers",
        headers=headers,
        json={
            "name": "Slot Customer",
            "code": "SLOT-CUSTOMER",
            "default_xlsx_template_id": template_response.json()["id"],
        },
    ).json()
    response = client.post(
        "/api/v1/generation-tasks",
        headers=headers,
        json={
            "name": "Too many products",
            "customer_id": customer["id"],
            "product_ids": product_ids,
            "output_template_version_id": version["id"],
        },
    )
    assert response.status_code == 400
    assert "仅配置了 1 个产品槽位" in response.json()["message"]


def test_pptx_renderer_preserves_native_shapes_and_replaces_image(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    destination = tmp_path / "generated.pptx"
    product_image = tmp_path / "product.png"
    template.write_bytes(_pptx_template())
    product_image.write_bytes(placeholder_png(320, 180))

    PptxRenderer().render(
        RenderRequest(
            template_path=template,
            destination=destination,
            mapping={
                "bindings": [
                    {
                        "object_key": "pptx:s1:shape:2",
                        "source": "product_name",
                        "product_slot": 1,
                        "visible": True,
                    },
                    {
                        "object_key": "pptx:s1:shape:3",
                        "source": "image.main",
                        "product_slot": 1,
                        "visible": True,
                        "fit": "contain",
                        "position": "center",
                        "fallback": ["unmatched_placeholder"],
                    },
                ]
            },
            products=[
                {
                    "fields": {"product_name": "Rendered Product"},
                    "images": [
                        {
                            "_path": str(product_image),
                            "image_type": "main",
                            "is_primary": True,
                            "sort_order": 0,
                            "width": 320,
                            "height": 180,
                            "mime_type": "image/png",
                        }
                    ],
                }
            ],
            customer={"name": "Customer", "currency": "USD", "settings": {}},
            parameters={},
        )
    )

    with ZipFile(destination) as archive:
        slide = archive.read("ppt/slides/slide1.xml").decode("utf-8")
        relationships = archive.read("ppt/slides/_rels/slide1.xml.rels").decode(
            "utf-8"
        )
        assert "Rendered Product" in slide
        assert "<ns0:pic>" in slide or "<p:pic>" in slide
        assert "rIdProductFlow3" in relationships
        assert "ppt/media/generated-s1-shape3.png" in archive.namelist()


def test_pptx_renderer_preserves_text_runs_and_selects_image_index(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    destination = tmp_path / "generated.pptx"
    first_image = tmp_path / "first.png"
    second_image = tmp_path / "second.png"
    template.write_bytes(_pptx_template())
    first_image.write_bytes(placeholder_png(320, 180))
    second_image.write_bytes(placeholder_png(120, 300))

    PptxRenderer().render(
        RenderRequest(
            template_path=template,
            destination=destination,
            mapping={
                "bindings": [
                    {
                        "object_key": "pptx:s1:shape:2",
                        "source": "product_name",
                        "product_slot": 1,
                        "visible": True,
                        "text_runs": [
                            {
                                "run_index": 0,
                                "source": "product_name",
                                "product_slot": 1,
                                "prefix": "Product: ",
                            },
                            {
                                "run_index": 1,
                                "source": "sku",
                                "visible": False,
                                "product_slot": 1,
                            },
                        ],
                    },
                    {
                        "object_key": "pptx:s1:shape:3",
                        "source": "image.main",
                        "product_slot": 1,
                        "visible": True,
                        "image_index": 1,
                        "fit": "contain",
                    },
                ]
            },
            products=[
                {
                    "fields": {"product_name": "Rendered Product", "sku": "RUN-001"},
                    "images": [
                        {
                            "_path": str(first_image),
                            "image_type": "main",
                            "sort_order": 0,
                            "width": 320,
                            "height": 180,
                            "mime_type": "image/png",
                        },
                        {
                            "_path": str(second_image),
                            "image_type": "main",
                            "sort_order": 1,
                            "width": 120,
                            "height": 300,
                            "mime_type": "image/png",
                        },
                    ],
                }
            ],
            customer={"name": "Customer", "currency": "USD", "settings": {}},
            parameters={},
        )
    )

    with ZipFile(destination) as archive:
        slide = archive.read("ppt/slides/slide1.xml").decode("utf-8")
        assert "Product: Rendered Product" in slide
        assert "Old suffix" not in slide
        assert 'x="1400100"' in slide
        assert 'cx="1200000"' in slide
        assert 'cy="3000000"' in slide
        with Image.open(BytesIO(archive.read("ppt/media/generated-s1-shape3.png"))) as image:
            assert image.size == (120, 300)
