import re
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image

from app.core.errors import BadRequestError
from app.services.renderers.base import (
    RenderRequest,
    format_text,
    resolve_image,
    resolve_value,
)


def _safe_cell_value(binding: dict, value, customer: dict, parameters: dict):
    formatter = binding.get("formatter") or "text"
    label = binding.get("label")
    if label or formatter in {"text", "date"}:
        return format_text(binding, value, customer=customer, parameters=parameters)
    if value is None or value == "":
        return ""
    try:
        number = float(value)
    except (TypeError, ValueError):
        return str(value)
    if formatter == "integer":
        return int(round(number))
    return number


def _prepare_image(payload: bytes, width: int, height: int, fit: str, position: str) -> BytesIO:
    with Image.open(BytesIO(payload)) as source:
        image = source.convert("RGBA")
        if fit == "stretch":
            result = image.resize((width, height))
        else:
            source_ratio = image.width / image.height
            target_ratio = width / height
            if fit == "cover":
                if source_ratio > target_ratio:
                    crop_width = round(image.height * target_ratio)
                    extra = image.width - crop_width
                    left = 0 if position == "left" else extra if position == "right" else extra // 2
                    image = image.crop((left, 0, left + crop_width, image.height))
                elif source_ratio < target_ratio:
                    crop_height = round(image.width / target_ratio)
                    extra = image.height - crop_height
                    top = 0 if position == "top" else extra if position == "bottom" else extra // 2
                    image = image.crop((0, top, image.width, top + crop_height))
                result = image.resize((width, height))
            else:
                image.thumbnail((width, height))
                result = Image.new("RGBA", (width, height), (255, 255, 255, 0))
                x = (
                    0
                    if position == "left"
                    else width - image.width
                    if position == "right"
                    else (width - image.width) // 2
                )
                y = (
                    0
                    if position == "top"
                    else height - image.height
                    if position == "bottom"
                    else (height - image.height) // 2
                )
                result.alpha_composite(image, (x, y))
        output = BytesIO()
        result.save(output, format="PNG")
        output.seek(0)
        return output


class XlsxRenderer:
    output_type = "xlsx"

    def render(self, request: RenderRequest) -> None:
        try:
            workbook = load_workbook(
                request.template_path, read_only=False, data_only=False, keep_links=False
            )
        except Exception as exc:
            raise BadRequestError("XLSX 模板无法读取。") from exc
        try:
            image_bindings: list[tuple[dict, int, int]] = []
            for binding in request.mapping.get("bindings", []):
                if not binding.get("visible", True):
                    continue
                cell_match = re.fullmatch(
                    r"xlsx:s(\d+):cell:([A-Z]+\d+)", binding["object_key"]
                )
                image_match = re.fullmatch(
                    r"xlsx:s(\d+):image:(\d+)", binding["object_key"]
                )
                if cell_match:
                    sheet_index = int(cell_match.group(1))
                    if sheet_index < 1 or sheet_index > len(workbook.worksheets):
                        raise BadRequestError(
                            "模板文件已经发生变化，请重新验证模板后再生成。"
                        )
                    cell = workbook.worksheets[sheet_index - 1][cell_match.group(2)]
                    value = resolve_value(binding, request.products, request.customer)
                    rendered = _safe_cell_value(
                        binding, value, request.customer, request.parameters
                    )
                    cell.value = rendered
                    if (
                        isinstance(rendered, str)
                        and rendered.startswith(("=", "+", "-", "@"))
                        and not binding.get("allow_formula", False)
                    ):
                        cell.data_type = "s"
                    formatter = binding.get("formatter")
                    if not binding.get("label") and formatter == "currency":
                        cell.number_format = '"$"#,##0.00'
                    elif not binding.get("label") and formatter == "percent":
                        cell.number_format = "0.0%"
                    elif not binding.get("label") and formatter == "integer":
                        cell.number_format = "#,##0"
                elif image_match:
                    image_bindings.append(
                        (binding, int(image_match.group(1)), int(image_match.group(2)))
                    )

            for binding, sheet_index, image_index in image_bindings:
                if sheet_index < 1 or sheet_index > len(workbook.worksheets):
                    raise BadRequestError(
                        "模板文件已经发生变化，请重新验证模板后再生成。"
                    )
                worksheet = workbook.worksheets[sheet_index - 1]
                if image_index < 1 or image_index > len(worksheet._images):
                    raise BadRequestError(
                        "模板文件已经发生变化，请重新验证模板后再生成。"
                    )
                original = worksheet._images[image_index - 1]
                width = max(1, round(original.width))
                height = max(1, round(original.height))
                image = resolve_image(
                    binding,
                    request.products,
                    request.customer,
                    placeholder_width=width,
                    placeholder_height=height,
                )
                prepared = _prepare_image(
                    image.payload,
                    width,
                    height,
                    binding.get("fit") or "contain",
                    binding.get("position") or "center",
                )
                replacement = ExcelImage(prepared)
                replacement.width = width
                replacement.height = height
                replacement.anchor = original.anchor
                worksheet._images[image_index - 1] = replacement
            request.destination.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(request.destination)
        finally:
            workbook.close()
        self.validate(request.destination)

    @staticmethod
    def validate(path: Path) -> None:
        try:
            workbook = load_workbook(path, read_only=False, data_only=False, keep_links=False)
            if not workbook.worksheets:
                raise BadRequestError("生成的 XLSX 不包含工作表。")
            workbook.close()
        except BadRequestError:
            raise
        except Exception as exc:
            raise BadRequestError("生成的 XLSX 未通过结构验证。") from exc
