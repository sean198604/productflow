from __future__ import annotations

import re
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from openpyxl import load_workbook

from app.core.errors import BadRequestError

PRESENTATION_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS = {"p": PRESENTATION_NS, "a": DRAWING_NS}
EMU_PER_POINT = 12700
MAX_ARCHIVE_FILES = 20_000
MAX_ARCHIVE_EXPANDED_BYTES = 512 * 1024 * 1024
MAX_XLSX_SCANNED_CELLS = 20_000
MAX_TEMPLATE_OBJECTS = 1_000


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _points(value: str | None) -> float | None:
    if not value:
        return None
    return round(int(value) / EMU_PER_POINT, 2)


def _serializable_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def validate_office_archive(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_FILES:
                raise BadRequestError("模板文件包含过多内部对象。")
            total_expanded = 0
            for member in members:
                if member.flag_bits & 0x1:
                    raise BadRequestError("模板文件已加密，请解除保护后重新上传。")
                total_expanded += member.file_size
                if total_expanded > MAX_ARCHIVE_EXPANDED_BYTES:
                    raise BadRequestError("模板解压后的内容过大。")
                if member.compress_size and member.file_size / member.compress_size > 1_000:
                    raise BadRequestError("模板压缩比例异常，无法安全解析。")
    except BadRequestError:
        raise
    except (OSError, zipfile.BadZipFile) as exc:
        raise BadRequestError("模板文件无法读取或不是有效的 Office 文件。") from exc


class TemplateParser:
    def parse(self, path: Path, output_type: str) -> dict[str, Any]:
        validate_office_archive(path)
        if output_type == "pptx":
            return self._parse_pptx(path)
        if output_type == "xlsx":
            return self._parse_xlsx(path)
        raise BadRequestError("仅支持 PPTX 和 XLSX 输出模板。")

    @staticmethod
    def _ppt_shape_type(element: ElementTree.Element) -> tuple[str, bool, str | None]:
        kind = _local_name(element.tag)
        if kind == "sp":
            if element.find("p:txBody", NS) is not None:
                return "text", True, None
            return "shape", False, "普通图形暂不支持字段参数化"
        if kind == "pic":
            return "image", True, None
        if kind == "graphicFrame":
            data = element.find(".//a:graphicData", NS)
            uri = data.get("uri", "") if data is not None else ""
            if uri.endswith("/table"):
                return "table", True, None
            if uri.endswith("/chart"):
                return "chart", False, "复杂图表暂不支持"
            if "diagram" in uri:
                return "smartart", False, "SmartArt 暂不支持"
            return "graphic_frame", False, "该图形对象暂不支持"
        if kind == "grpSp":
            return "group", False, "组合图形暂不支持"
        if kind == "cxnSp":
            return "connector", False, "连接线暂不支持"
        if kind == "contentPart":
            return "content_part", False, "外部内容对象暂不支持"
        return kind or "unknown", False, "该 PowerPoint 对象暂不支持"

    @staticmethod
    def _ppt_position(element: ElementTree.Element) -> tuple[dict, dict]:
        offset = element.find(".//a:off", NS)
        extent = element.find(".//a:ext", NS)
        return (
            {
                "x": _points(offset.get("x")) if offset is not None else None,
                "y": _points(offset.get("y")) if offset is not None else None,
            },
            {
                "width": _points(extent.get("cx")) if extent is not None else None,
                "height": _points(extent.get("cy")) if extent is not None else None,
            },
        )

    @staticmethod
    def _ppt_font(element: ElementTree.Element) -> dict[str, Any]:
        run_properties = element.find(".//a:rPr", NS)
        if run_properties is None:
            run_properties = element.find(".//a:defRPr", NS)
        if run_properties is None:
            return {}
        latin = run_properties.find("a:latin", NS)
        color = run_properties.find(".//a:srgbClr", NS)
        size = run_properties.get("sz")
        return {
            "name": latin.get("typeface") if latin is not None else None,
            "size": round(int(size) / 100, 2) if size else None,
            "bold": run_properties.get("b") == "1",
            "italic": run_properties.get("i") == "1",
            "color": f"#{color.get('val')}" if color is not None else None,
        }

    def _parse_pptx(self, path: Path) -> dict[str, Any]:
        try:
            with zipfile.ZipFile(path) as archive:
                names = set(archive.namelist())
                required = {"[Content_Types].xml", "ppt/presentation.xml"}
                if not required.issubset(names):
                    raise BadRequestError("PPTX 模板缺少必要的演示文稿结构。")
                presentation = ElementTree.fromstring(archive.read("ppt/presentation.xml"))
                slide_size = presentation.find("p:sldSz", NS)
                slide_width = _points(slide_size.get("cx")) if slide_size is not None else None
                slide_height = _points(slide_size.get("cy")) if slide_size is not None else None
                slide_files = sorted(
                    (
                        name
                        for name in names
                        if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
                    ),
                    key=lambda name: int(re.search(r"(\d+)", Path(name).stem).group(1)),
                )
                if not slide_files:
                    raise BadRequestError("PPTX 模板中没有可用的幻灯片。")

                objects: list[dict[str, Any]] = []
                warnings: list[str] = []
                for slide_index, slide_file in enumerate(slide_files, start=1):
                    root = ElementTree.fromstring(archive.read(slide_file))
                    objects.append(
                        {
                            "object_key": f"pptx:s{slide_index}:slide",
                            "container": f"Slide {slide_index}",
                            "slide_index": slide_index,
                            "shape_id": None,
                            "shape_name": f"Slide {slide_index}",
                            "object_type": "slide",
                            "supported": True,
                            "reason": None,
                            "preview": None,
                            "position": {"x": 0, "y": 0},
                            "size": {"width": slide_width, "height": slide_height},
                            "font": {},
                        }
                    )
                    if root.find("p:timing", NS) is not None:
                        warnings.append(f"Slide {slide_index} 包含动画，生成时不会修改动画。")
                    shape_tree = root.find(".//p:cSld/p:spTree", NS)
                    if shape_tree is None:
                        continue
                    for element in list(shape_tree):
                        if _local_name(element.tag) in {"nvGrpSpPr", "grpSpPr"}:
                            continue
                        c_nv_pr = element.find(".//p:cNvPr", NS)
                        shape_id = int(c_nv_pr.get("id", "0")) if c_nv_pr is not None else 0
                        shape_name = (
                            c_nv_pr.get("name", f"Shape {shape_id}")
                            if c_nv_pr is not None
                            else f"Shape {shape_id}"
                        )
                        object_type, supported, reason = self._ppt_shape_type(element)
                        position, size = self._ppt_position(element)
                        text_parts = [node.text or "" for node in element.findall(".//a:t", NS)]
                        preview = " ".join(part.strip() for part in text_parts if part.strip())
                        objects.append(
                            {
                                "object_key": f"pptx:s{slide_index}:shape:{shape_id}",
                                "container": f"Slide {slide_index}",
                                "slide_index": slide_index,
                                "shape_id": shape_id,
                                "shape_name": shape_name,
                                "object_type": object_type,
                                "supported": supported,
                                "reason": reason,
                                "preview": preview[:300] or None,
                                "position": position,
                                "size": size,
                                "font": self._ppt_font(element),
                            }
                        )
        except BadRequestError:
            raise
        except (KeyError, ElementTree.ParseError, ValueError) as exc:
            raise BadRequestError("PPTX 模板结构损坏，无法完成验证。") from exc

        supported_count = sum(item["supported"] for item in objects)
        unsupported_count = len(objects) - supported_count
        return {
            "format": "pptx",
            "is_valid": True,
            "object_count": len(objects),
            "parameterizable_count": supported_count,
            "unsupported_count": unsupported_count,
            "objects": objects,
            "warnings": sorted(set(warnings)),
            "errors": [],
            "capabilities": [
                "slide",
                "text",
                "image",
                "table",
                "position",
                "size",
                "basic_font",
            ],
            "unsupported_capabilities": [
                "smartart",
                "complex_chart",
                "animation_editing",
                "vba",
                "complex_master_editing",
            ],
        }

    @staticmethod
    def _xlsx_font(cell) -> dict[str, Any]:
        color = None
        if cell.font.color is not None and cell.font.color.type == "rgb":
            color = cell.font.color.rgb
            if color and len(color) == 8:
                color = color[2:]
        return {
            "name": cell.font.name,
            "size": float(cell.font.sz) if cell.font.sz is not None else None,
            "bold": bool(cell.font.b),
            "italic": bool(cell.font.i),
            "color": f"#{color}" if color else None,
            "number_format": cell.number_format,
        }

    def _parse_xlsx(self, path: Path) -> dict[str, Any]:
        try:
            workbook = load_workbook(path, read_only=False, data_only=False, keep_links=False)
        except Exception as exc:
            raise BadRequestError("XLSX 模板结构损坏，无法完成验证。") from exc
        objects: list[dict[str, Any]] = []
        warnings: list[str] = []
        scanned_cells = 0
        try:
            # Collect workbook structure and drawings across every sheet before
            # scanning cells. Large formatted ranges must not hide image slots.
            for sheet_index, worksheet in enumerate(workbook.worksheets, start=1):
                if len(objects) >= MAX_TEMPLATE_OBJECTS:
                    break
                objects.append(
                    {
                        "object_key": f"xlsx:s{sheet_index}:sheet",
                        "container": worksheet.title,
                        "sheet_name": worksheet.title,
                        "object_type": "sheet",
                        "supported": True,
                        "reason": None,
                        "preview": None,
                        "position": {"row": 1, "column": 1},
                        "size": {"max_row": worksheet.max_row, "max_column": worksheet.max_column},
                        "font": {},
                    }
                )
                for image_index, image in enumerate(worksheet._images, start=1):
                    if len(objects) >= MAX_TEMPLATE_OBJECTS:
                        break
                    anchor = getattr(image.anchor, "_from", None)
                    row_number = anchor.row + 1 if anchor is not None else None
                    column_number = anchor.col + 1 if anchor is not None else None
                    objects.append(
                        {
                            "object_key": f"xlsx:s{sheet_index}:image:{image_index}",
                            "container": worksheet.title,
                            "sheet_name": worksheet.title,
                            "object_type": "image",
                            "supported": True,
                            "reason": None,
                            "preview": getattr(image, "path", None),
                            "position": {"row": row_number, "column": column_number},
                            "size": {"width": image.width, "height": image.height},
                            "font": {},
                        }
                    )

                for table in worksheet.tables.values():
                    if len(objects) >= MAX_TEMPLATE_OBJECTS:
                        break
                    objects.append(
                        {
                            "object_key": f"xlsx:s{sheet_index}:table:{table.name}",
                            "container": worksheet.title,
                            "sheet_name": worksheet.title,
                            "object_type": "table",
                            "supported": True,
                            "reason": None,
                            "preview": table.name,
                            "position": {"range": table.ref},
                            "size": {},
                            "font": {},
                        }
                    )

                for chart_index, _chart in enumerate(worksheet._charts, start=1):
                    if len(objects) >= MAX_TEMPLATE_OBJECTS:
                        break
                    objects.append(
                        {
                            "object_key": f"xlsx:s{sheet_index}:chart:{chart_index}",
                            "container": worksheet.title,
                            "sheet_name": worksheet.title,
                            "object_type": "chart",
                            "supported": False,
                            "reason": "复杂图表暂不支持参数化",
                            "preview": None,
                            "position": {},
                            "size": {},
                            "font": {},
                        }
                    )

            for sheet_index, worksheet in enumerate(workbook.worksheets, start=1):
                if len(objects) >= MAX_TEMPLATE_OBJECTS:
                    warnings.append(
                        "模板对象较多，当前版本仅展示前 1,000 个可验证对象。"
                    )
                    break
                merged_ranges = list(worksheet.merged_cells.ranges)
                for row in worksheet.iter_rows():
                    for cell in row:
                        scanned_cells += 1
                        if scanned_cells > MAX_XLSX_SCANNED_CELLS:
                            warnings.append(
                                "工作簿使用区域较大，当前版本仅分析前 20,000 个单元格。"
                            )
                            break
                        if len(objects) >= MAX_TEMPLATE_OBJECTS:
                            warnings.append(
                                "模板对象较多，当前版本仅展示前 1,000 个可验证对象。"
                            )
                            break
                        if cell.value is None and cell.style_id == 0:
                            continue
                        merged_range = next(
                            (item for item in merged_ranges if cell.coordinate in item), None
                        )
                        if (
                            merged_range is not None
                            and cell.coordinate != merged_range.start_cell.coordinate
                        ):
                            continue
                        is_formula = cell.data_type == "f"
                        object_type = "merged_cell" if merged_range is not None else "cell"
                        objects.append(
                            {
                                "object_key": f"xlsx:s{sheet_index}:cell:{cell.coordinate}",
                                "container": worksheet.title,
                                "sheet_name": worksheet.title,
                                "cell": cell.coordinate,
                                "object_type": object_type,
                                "supported": not is_formula,
                                "reason": "公式单元格不可直接绑定字段" if is_formula else None,
                                "preview": _serializable_value(cell.value),
                                "position": {"row": cell.row, "column": cell.column},
                                "size": {
                                    "merged_range": str(merged_range) if merged_range else None
                                },
                                "font": self._xlsx_font(cell),
                            }
                        )
                    if (
                        scanned_cells > MAX_XLSX_SCANNED_CELLS
                        or len(objects) >= MAX_TEMPLATE_OBJECTS
                    ):
                        break
                if len(objects) >= MAX_TEMPLATE_OBJECTS:
                    break

        finally:
            workbook.close()

        supported_count = sum(item["supported"] for item in objects)
        unsupported_count = len(objects) - supported_count
        return {
            "format": "xlsx",
            "is_valid": True,
            "object_count": len(objects),
            "parameterizable_count": supported_count,
            "unsupported_count": unsupported_count,
            "objects": objects,
            "warnings": sorted(set(warnings)),
            "errors": [],
            "capabilities": [
                "sheet",
                "cell",
                "merged_cell",
                "image",
                "table",
                "position",
                "size",
                "basic_font",
            ],
            "unsupported_capabilities": [
                "complex_chart",
                "vba",
                "arbitrary_excel_object_editing",
            ],
        }
